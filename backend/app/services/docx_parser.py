
from pathlib import Path

from docx import Document


def extract_text_from_docx(file_path: str) -> str:
    """
    Extract text from paragraphs and tables in a DOCX file.

    Returns an empty string when the document contains no extractable text.
    Raises FileNotFoundError when the file does not exist.
    Document parsing errors are allowed to propagate to the caller.
    """

    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    document = Document(str(path))
    extracted_sections: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            extracted_sections.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip().replace("\n", " ")
                for cell in row.cells
            ]
            non_empty_cells = [cell for cell in cells if cell]

            if non_empty_cells:
                extracted_sections.append(" | ".join(non_empty_cells))

    return "\n".join(extracted_sections).strip()
