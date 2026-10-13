
from pathlib import Path

import pytest
from docx import Document

from app.services.docx_parser import extract_text_from_docx


def test_extracts_paragraph_text(tmp_path: Path) -> None:
    file_path = tmp_path / "paragraph_resume.docx"

    document = Document()
    document.add_paragraph("Python developer")
    document.add_paragraph("Experience: 3 years")
    document.save(file_path)

    result = extract_text_from_docx(str(file_path))

    assert "Python developer" in result
    assert "Experience: 3 years" in result


def test_extracts_table_text(tmp_path: Path) -> None:
    file_path = tmp_path / "table_resume.docx"

    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "Python"
    table.cell(1, 0).text = "Experience"
    table.cell(1, 1).text = "3 years"
    document.save(file_path)

    result = extract_text_from_docx(str(file_path))

    assert "Skill | Python" in result
    assert "Experience | 3 years" in result


def test_empty_document_returns_empty_string(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "empty_resume.docx"

    Document().save(file_path)

    assert extract_text_from_docx(str(file_path)) == ""


def test_missing_file_raises_file_not_found(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "missing.docx"

    with pytest.raises(FileNotFoundError):
        extract_text_from_docx(str(file_path))


def test_corrupt_docx_raises_parsing_error(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "corrupt.docx"
    file_path.write_text("This is not a DOCX file.", encoding="utf-8")

    with pytest.raises(Exception):
        extract_text_from_docx(str(file_path))
